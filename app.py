import csv
import io
import os
import time
from collections import defaultdict, deque
from pathlib import Path

from docx import Document
from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import HTMLResponse, Response
from openpyxl import load_workbook
from pydantic import BaseModel
from pypdf import PdfReader
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

from agent_workflow import run_assurance as run_consumer_duty_assurance
from motor_finance_workflow import run_assurance as run_motor_finance_assurance
from remediation_workflow import run_assurance as run_remediation_assurance


app = FastAPI(title="Agentic Financial Services Assurance Portfolio")

MAX_FIELD_CHARS = 12000
RATE_LIMIT = 8
RATE_WINDOW_SECONDS = 3600
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_UPLOAD_FILES = 3
MAX_EXTRACTED_CHARS_PER_FILE = 18000
ALLOWED_UPLOAD_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv", ".xlsx"}
_requests_by_ip: dict[str, deque[float]] = defaultdict(deque)

# Lightweight portfolio usage counters. These are intentionally in-memory:
# they reset whenever Render restarts or redeploys the service.
_service_started_at = time.time()
_demo_usage = {
    "consumer-duty": {"successful_runs": 0, "failed_runs": 0},
    "motor-finance": {"successful_runs": 0, "failed_runs": 0},
    "remediation": {"successful_runs": 0, "failed_runs": 0},
}


def _record_demo_run(demo: str, success: bool) -> None:
    if demo not in _demo_usage:
        return
    key = "successful_runs" if success else "failed_runs"
    _demo_usage[demo][key] += 1


class ConsumerDutyCaseInput(BaseModel):
    customer_circumstances: str
    interaction_notes: str
    vulnerability_support_needs: str
    actions_taken: str
    agent_rationale: str
    uploaded_evidence: str = ""


class MotorFinanceCaseInput(BaseModel):
    agreement_details: str
    commission_evidence: str
    arrangement_classification: str
    disclosure_customer_evidence: str
    proposed_outcome: str
    reviewer_rationale: str
    uploaded_evidence: str = ""


class RemediationProgrammeInput(BaseModel):
    customer_harm: str
    population_identification: str
    data_lineage: str
    redress_methodology: str
    qa_outcome_testing: str
    governance_closure: str
    uploaded_evidence: str = ""


def _client_ip(request: Request) -> str:
    forwarded = request.headers.get("x-forwarded-for")
    if forwarded:
        return forwarded.split(",")[0].strip()
    return request.client.host if request.client else "unknown"


def _check_rate_limit(request: Request) -> None:
    now = time.time()
    ip = _client_ip(request)
    bucket = _requests_by_ip[ip]
    while bucket and now - bucket[0] > RATE_WINDOW_SECONDS:
        bucket.popleft()
    if len(bucket) >= RATE_LIMIT:
        raise HTTPException(
            status_code=429,
            detail="Demo limit reached for this hour. Please try again later.",
        )
    bucket.append(now)


def _validate_demo_input(case) -> None:
    for field_name, value in case.model_dump().items():
        limit = MAX_EXTRACTED_CHARS_PER_FILE * MAX_UPLOAD_FILES + 2000 if field_name == "uploaded_evidence" else MAX_FIELD_CHARS
        if len(value) > limit:
            raise HTTPException(
                status_code=400,
                detail=f"{field_name} is too long for this portfolio demo.",
            )


def _require_api_key() -> None:
    if not os.getenv("OPENAI_API_KEY"):
        raise HTTPException(
            status_code=503,
            detail="OPENAI_API_KEY is not configured for the live demo.",
        )


def _extract_pdf(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    parts = []
    for index, page in enumerate(reader.pages[:40], start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"[Page {index}]\n{text.strip()}")
    return "\n\n".join(parts)


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table_index, table in enumerate(doc.tables[:20], start=1):
        paragraphs.append(f"[Table {table_index}]")
        for row in table.rows[:100]:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(paragraphs)


def _extract_txt_or_csv(data: bytes, ext: str) -> str:
    text = data.decode("utf-8", errors="replace")
    if ext == ".csv":
        reader = csv.reader(io.StringIO(text))
        rows = []
        for i, row in enumerate(reader):
            if i >= 500:
                break
            rows.append(" | ".join(row))
        return "\n".join(rows)
    return text


def _extract_xlsx(data: bytes) -> str:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts = []
    total_rows = 0
    for ws in wb.worksheets[:10]:
        parts.append(f"[Sheet: {ws.title}]")
        for row in ws.iter_rows(values_only=True):
            values = ["" if v is None else str(v) for v in row[:30]]
            if any(values):
                parts.append(" | ".join(values))
                total_rows += 1
            if total_rows >= 500:
                break
        if total_rows >= 500:
            break
    return "\n".join(parts)


def _extract_file_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_UPLOAD_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"{filename}: unsupported file type. Use PDF, DOCX, TXT, CSV or XLSX.",
        )
    try:
        if ext == ".pdf":
            text = _extract_pdf(data)
        elif ext == ".docx":
            text = _extract_docx(data)
        elif ext in {".txt", ".csv"}:
            text = _extract_txt_or_csv(data, ext)
        else:
            text = _extract_xlsx(data)
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=f"{filename}: the file could not be read ({exc}).",
        ) from exc

    text = text.strip()
    if not text:
        if ext == ".pdf":
            raise HTTPException(
                status_code=400,
                detail=f"{filename}: no extractable text was found. Image-only/scanned PDFs are not supported in this public demo.",
            )
        raise HTTPException(status_code=400, detail=f"{filename}: no readable text was found.")
    return text[:MAX_EXTRACTED_CHARS_PER_FILE]


def _safe_pdf_text(value) -> str:
    if value is None:
        return ""
    return str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _build_remediation_pdf(payload: dict) -> bytes:
    result = payload.get("result") or {}
    human_decision = payload.get("human_decision") or "Not recorded"
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title="Agentic Redress & Remediation Assurance Report",
        author="Portfolio demonstration",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="TitleCenter", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="SmallMuted", parent=styles["BodyText"], fontSize=8, leading=10, textColor="#667085"))
    styles.add(ParagraphStyle(name="Section", parent=styles["Heading2"], spaceBefore=10, spaceAfter=6))
    story = [
        Paragraph("Agentic Redress &amp; Remediation Assurance", styles["TitleCenter"]),
        Paragraph("Portfolio demonstration · Fictional / anonymised data only", styles["SmallMuted"]),
        Spacer(1, 8),
        Paragraph(f"<b>Recommendation:</b> {_safe_pdf_text(result.get('recommendation', '—'))}", styles["Heading2"]),
        Paragraph(_safe_pdf_text(result.get("programme_summary", "")), styles["BodyText"]),
        Spacer(1, 5),
        Paragraph(_safe_pdf_text(result.get("rationale", "")), styles["BodyText"]),
    ]

    drivers = result.get("decision_drivers") or []
    if drivers:
        story.append(Paragraph("Decision drivers", styles["Section"]))
        for item in drivers:
            story.append(Paragraph("• " + _safe_pdf_text(item), styles["BodyText"]))

    risk = result.get("risk_dashboard") or []
    if risk:
        story.append(Paragraph("Assurance risk dashboard", styles["Section"]))
        data = [["Area", "Risk", "Reason"]]
        for row in risk:
            data.append([
                _safe_pdf_text(row.get("area", "")),
                _safe_pdf_text(row.get("rating", "")),
                Paragraph(_safe_pdf_text(row.get("reason", "")), styles["BodyText"]),
            ])
        table = Table(data, colWidths=[42 * mm, 22 * mm, 100 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), "#EEF3FF"),
            ("GRID", (0, 0), (-1, -1), 0.4, "#C8D2E2"),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("LEADING", (0, 0), (-1, -1), 10),
        ]))
        story.append(table)

    rca = result.get("root_cause_analysis") or {}
    if rca:
        story.append(Paragraph("RCA flow", styles["Section"]))
        flow_labels = [
            ("Root cause", rca.get("root_cause")),
            ("Customer harm", rca.get("customer_harm")),
            ("Population risk", rca.get("population_risk")),
            ("Data issue", rca.get("data_issue")),
            ("Methodology impact", rca.get("methodology_impact")),
            ("QA issue", rca.get("qa_issue")),
            ("Closure risk", rca.get("closure_risk")),
        ]
        for label, value in flow_labels:
            if value:
                story.append(Paragraph(f"<b>{_safe_pdf_text(label)}:</b> {_safe_pdf_text(value)}", styles["BodyText"]))

    findings = result.get("findings") or []
    if findings:
        story.append(PageBreak())
        story.append(Paragraph("Detailed findings", styles["Section"]))
        for i, finding in enumerate(findings, start=1):
            story.append(Paragraph(
                f"<b>{i}. {_safe_pdf_text(finding.get('area',''))} · {_safe_pdf_text(str(finding.get('severity','')).upper())}</b>",
                styles["Heading3"],
            ))
            story.append(Paragraph("<b>Issue:</b> " + _safe_pdf_text(finding.get("issue", "")), styles["BodyText"]))
            story.append(Paragraph("<b>Why it matters:</b> " + _safe_pdf_text(finding.get("why_it_matters", "")), styles["BodyText"]))
            story.append(Paragraph("<b>Reviewer action:</b> " + _safe_pdf_text(finding.get("reviewer_action", "")), styles["BodyText"]))
            for ev in finding.get("evidence_refs") or []:
                story.append(Paragraph(
                    f"<b>Evidence:</b> [{_safe_pdf_text(ev.get('source',''))}] {_safe_pdf_text(ev.get('evidence',''))}",
                    styles["SmallMuted"],
                ))
            if finding.get("regulatory_reference"):
                story.append(Paragraph(
                    "<b>Regulatory reference to verify:</b> " + _safe_pdf_text(finding.get("regulatory_reference")),
                    styles["SmallMuted"],
                ))
            story.append(Spacer(1, 6))

    evidence = result.get("evidence_to_obtain") or []
    if evidence:
        story.append(Paragraph("Evidence to obtain / verify", styles["Section"]))
        for item in evidence:
            story.append(Paragraph("• " + _safe_pdf_text(item), styles["BodyText"]))

    story.append(Paragraph("Human decision", styles["Section"]))
    story.append(Paragraph(_safe_pdf_text(human_decision), styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "This report is generated by a portfolio prototype. It does not make legal, regulatory, methodology-approval or compensation decisions. Human review remains required.",
        styles["SmallMuted"],
    ))
    doc.build(story)
    return buffer.getvalue()



def _build_case_assurance_pdf(payload: dict, title: str, mode: str) -> bytes:
    result = payload.get("result") or {}
    human_decision = payload.get("human_decision") or "Not recorded"
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=title,
        author="Portfolio demonstration",
    )
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CaseTitleCenter", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="CaseSmallMuted", parent=styles["BodyText"], fontSize=8, leading=10, textColor="#667085"))
    styles.add(ParagraphStyle(name="CaseSection", parent=styles["Heading2"], spaceBefore=10, spaceAfter=6))

    story = [
        Paragraph(_safe_pdf_text(title), styles["CaseTitleCenter"]),
        Paragraph("Portfolio demonstration · Fictional / anonymised data only", styles["CaseSmallMuted"]),
        Spacer(1, 8),
        Paragraph(f"<b>Recommendation:</b> {_safe_pdf_text(result.get('recommendation', '—'))}", styles["Heading2"]),
        Paragraph(_safe_pdf_text(result.get("case_summary", "")), styles["BodyText"]),
        Spacer(1, 5),
        Paragraph(_safe_pdf_text(result.get("rationale", "")), styles["BodyText"]),
    ]

    drivers = result.get("decision_drivers") or result.get("escalation_drivers") or []
    if drivers:
        story.append(Paragraph("Decision drivers", styles["CaseSection"]))
        for item in drivers:
            story.append(Paragraph("• " + _safe_pdf_text(item), styles["BodyText"]))

    if mode == "motor":
        matrix = result.get("evidence_matrix") or []
        if matrix:
            story.append(Paragraph("Evidence sufficiency matrix", styles["CaseSection"]))
            rows = [["Evidence required", "Status", "Source", "Gap / observation"]]
            for item in matrix:
                rows.append([
                    Paragraph(_safe_pdf_text(item.get("evidence_required", "")), styles["BodyText"]),
                    _safe_pdf_text(item.get("status", "")),
                    Paragraph(_safe_pdf_text(item.get("source", "")), styles["BodyText"]),
                    Paragraph(_safe_pdf_text(item.get("gap_or_observation", "")), styles["BodyText"]),
                ])
            table = Table(rows, colWidths=[42 * mm, 22 * mm, 43 * mm, 57 * mm], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), "#EEF3FF"),
                ("GRID", (0, 0), (-1, -1), 0.4, "#C8D2E2"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEADING", (0, 0), (-1, -1), 9),
            ]))
            story.append(table)

        flow = result.get("agreement_evidence_flow") or {}
        if flow:
            story.append(Paragraph("Agreement evidence flow", styles["CaseSection"]))
            for label, key in [
                ("Agreement", "agreement"),
                ("Commission evidence", "commission_evidence"),
                ("Arrangement type", "arrangement_type"),
                ("Disclosure", "disclosure"),
                ("Customer evidence", "customer_evidence"),
                ("Scheme pathway", "scheme_pathway"),
                ("Assurance outcome", "assurance_outcome"),
            ]:
                story.append(Paragraph(f"<b>{label}:</b> {_safe_pdf_text(flow.get(key, ''))}", styles["BodyText"]))

    if mode == "consumer":
        journey = result.get("customer_journey") or []
        if journey:
            story.append(Paragraph("Customer journey reconstruction", styles["CaseSection"]))
            rows = [["Stage", "Status", "Event / evidence", "Source"]]
            for item in journey:
                rows.append([
                    Paragraph(_safe_pdf_text(item.get("stage", "")), styles["BodyText"]),
                    _safe_pdf_text(item.get("status", "")),
                    Paragraph(_safe_pdf_text(item.get("event", "")), styles["BodyText"]),
                    Paragraph(_safe_pdf_text(item.get("evidence_source", "")), styles["BodyText"]),
                ])
            table = Table(rows, colWidths=[34 * mm, 24 * mm, 72 * mm, 34 * mm], repeatRows=1)
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), "#EEF3FF"),
                ("GRID", (0, 0), (-1, -1), 0.4, "#C8D2E2"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("LEADING", (0, 0), (-1, -1), 9),
            ]))
            story.append(table)

    findings = result.get("findings") or []
    if findings:
        story.append(PageBreak())
        story.append(Paragraph("Detailed findings", styles["CaseSection"]))
        for i, finding in enumerate(findings, start=1):
            story.append(Paragraph(
                f"<b>{i}. {_safe_pdf_text(finding.get('area',''))} · {_safe_pdf_text(str(finding.get('severity','')).upper())}</b>",
                styles["Heading3"],
            ))
            story.append(Paragraph("<b>Issue:</b> " + _safe_pdf_text(finding.get("issue", "")), styles["BodyText"]))
            story.append(Paragraph("<b>Why it matters:</b> " + _safe_pdf_text(finding.get("why_it_matters", "")), styles["BodyText"]))
            story.append(Paragraph("<b>Reviewer action:</b> " + _safe_pdf_text(finding.get("reviewer_action", "")), styles["BodyText"]))
            for ev in finding.get("evidence_refs") or []:
                story.append(Paragraph(
                    f"<b>Evidence:</b> [{_safe_pdf_text(ev.get('source',''))}] {_safe_pdf_text(ev.get('evidence',''))}",
                    styles["CaseSmallMuted"],
                ))
            if finding.get("fca_reference"):
                story.append(Paragraph(
                    "<b>FCA reference to verify:</b> " + _safe_pdf_text(finding.get("fca_reference")),
                    styles["CaseSmallMuted"],
                ))
            story.append(Spacer(1, 6))

    evidence = result.get("evidence_to_obtain") or []
    if evidence:
        story.append(Paragraph("Evidence to obtain / verify", styles["CaseSection"]))
        for item in evidence:
            story.append(Paragraph("• " + _safe_pdf_text(item), styles["BodyText"]))

    story.append(Paragraph("Human decision", styles["CaseSection"]))
    story.append(Paragraph(_safe_pdf_text(human_decision), styles["BodyText"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "This report is generated by a portfolio prototype. It does not make legal, regulatory, eligibility or compensation decisions. Human review remains required.",
        styles["CaseSmallMuted"],
    ))
    doc.build(story)
    return buffer.getvalue()


async def _extract_uploaded_files(files: list[UploadFile]) -> dict:
    if not files or len(files) > MAX_UPLOAD_FILES:
        raise HTTPException(status_code=400, detail=f"Upload between 1 and {MAX_UPLOAD_FILES} files.")
    extracted = []
    for upload in files:
        filename = Path(upload.filename or "upload").name
        data = await upload.read(MAX_UPLOAD_BYTES + 1)
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(status_code=400, detail=f"{filename}: file exceeds the 2 MB demo limit.")
        text = _extract_file_text(filename, data)
        extracted.append({"filename": filename, "text": text, "characters": len(text)})
    return {"files": extracted}


@app.get("/", response_class=HTMLResponse)
async def home():
    with open("index.html", "r", encoding="utf-8") as f:
        return f.read()


@app.get("/consumer-duty", response_class=HTMLResponse)
async def consumer_duty_page():
    with open("consumer_duty.html", "r", encoding="utf-8") as f:
        return f.read()


@app.get("/motor-finance", response_class=HTMLResponse)
async def motor_finance_page():
    with open("motor_finance.html", "r", encoding="utf-8") as f:
        return f.read()


@app.get("/remediation", response_class=HTMLResponse)
async def remediation_page():
    with open("remediation.html", "r", encoding="utf-8") as f:
        return f.read()


@app.get("/api/demo-usage")
async def demo_usage():
    total_successful = sum(v["successful_runs"] for v in _demo_usage.values())
    total_failed = sum(v["failed_runs"] for v in _demo_usage.values())
    return {
        "api_configured": bool(os.getenv("OPENAI_API_KEY")),
        "service_started_at_unix": int(_service_started_at),
        "note": "Counters reset when the Render service restarts or redeploys.",
        "total_successful_runs": total_successful,
        "total_failed_runs": total_failed,
        "demos": _demo_usage,
    }


@app.get("/health")
async def health():
    return {
        "status": "ok",
        "agentic_mode": bool(os.getenv("OPENAI_API_KEY")),
        "demos": ["consumer-duty", "motor-finance", "remediation"],
        "remediation_features": ["document-extraction", "evidence-traceability", "risk-dashboard", "rca-flow", "pdf-report"],
        "motor_finance_features": ["agreement-evidence-upload", "evidence-matrix", "evidence-flow", "pdf-report"],
        "consumer_duty_features": ["journey-evidence-upload", "customer-journey", "evidence-traceability", "pdf-report"],
    }


@app.post("/api/consumer-duty/review")
async def consumer_duty_review(case: ConsumerDutyCaseInput, request: Request):
    _require_api_key()
    _check_rate_limit(request)
    _validate_demo_input(case)
    try:
        result = await run_consumer_duty_assurance(case.model_dump())
        _record_demo_run("consumer-duty", True)
        return result.model_dump()
    except HTTPException:
        _record_demo_run("consumer-duty", False)
        raise
    except Exception as exc:
        _record_demo_run("consumer-duty", False)
        raise HTTPException(status_code=500, detail=f"Agentic review failed: {exc}") from exc


@app.post("/api/motor-finance/review")
async def motor_finance_review(case: MotorFinanceCaseInput, request: Request):
    _require_api_key()
    _check_rate_limit(request)
    _validate_demo_input(case)
    try:
        result = await run_motor_finance_assurance(case.model_dump())
        _record_demo_run("motor-finance", True)
        return result.model_dump()
    except HTTPException:
        _record_demo_run("motor-finance", False)
        raise
    except Exception as exc:
        _record_demo_run("motor-finance", False)
        raise HTTPException(status_code=500, detail=f"Agentic review failed: {exc}") from exc


@app.post("/api/remediation/extract")
async def remediation_extract(request: Request, files: list[UploadFile] = File(...)):
    _check_rate_limit(request)
    return await _extract_uploaded_files(files)


@app.post("/api/motor-finance/extract")
async def motor_finance_extract(request: Request, files: list[UploadFile] = File(...)):
    _check_rate_limit(request)
    return await _extract_uploaded_files(files)


@app.post("/api/consumer-duty/extract")
async def consumer_duty_extract(request: Request, files: list[UploadFile] = File(...)):
    _check_rate_limit(request)
    return await _extract_uploaded_files(files)


@app.post("/api/remediation/review")
async def remediation_review(case: RemediationProgrammeInput, request: Request):
    _require_api_key()
    _check_rate_limit(request)
    _validate_demo_input(case)
    try:
        result = await run_remediation_assurance(case.model_dump())
        _record_demo_run("remediation", True)
        return result.model_dump()
    except HTTPException:
        _record_demo_run("remediation", False)
        raise
    except Exception as exc:
        _record_demo_run("remediation", False)
        raise HTTPException(status_code=500, detail=f"Agentic review failed: {exc}") from exc


@app.post("/api/remediation/report")
async def remediation_report(payload: dict):
    pdf_bytes = _build_remediation_pdf(payload)
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="remediation-assurance-report.pdf"'},
    )


@app.post("/api/motor-finance/report")
async def motor_finance_report(payload: dict):
    pdf_bytes = _build_case_assurance_pdf(
        payload,
        "Agentic Motor Finance Remediation Assurance Report",
        "motor",
    )
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="motor-finance-assurance-report.pdf"'},
    )


@app.post("/api/consumer-duty/report")
async def consumer_duty_report(payload: dict):
    pdf_bytes = _build_case_assurance_pdf(
        payload,
        "Agentic Consumer Duty Assurance Report",
        "consumer",
    )
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="consumer-duty-assurance-report.pdf"'},
    )


# Backwards compatibility for the original Consumer Duty demo endpoint.
@app.post("/api/review")
async def legacy_consumer_duty_review(case: ConsumerDutyCaseInput, request: Request):
    return await consumer_duty_review(case, request)
